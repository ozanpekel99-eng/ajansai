"""Sohbetten dosya üretimi: görselli sunum (PDF/PPTX), belge (PDF/DOCX) ve tablo (XLSX).

Akış: sohbet modeli [[DOSYA: tür|format|açıklama]] satırı yazar → burada ikinci bir model
çağrısıyla içerik JSON olarak hazırlanır → dosya Python kütüphaneleriyle oluşturulur.
Sunum slaytlarına Pixabay'den konuya uygun telifsiz fotoğraf eklenir (ücretsiz).
"""

import io
import json
import logging
import re
from datetime import date
from pathlib import Path

import requests

from app.services.llm_service import LLMError, smart_chat_completion

logger = logging.getLogger("ajansai.documents")

FILE_TAG = re.compile(r"\[\[\s*DOSYA\s*:\s*(.+?)\]\]", re.IGNORECASE | re.DOTALL)
FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"

KINDS = {"sunum": "presentation", "belge": "document", "tablo": "spreadsheet"}
FORMATS = {
    "presentation": ("pdf", "pptx"),
    "document": ("pdf", "docx"),
    "spreadsheet": ("xlsx",),
}
MIME = {
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

FILE_NOTE = (
    "\n- Dosya oluşturabilirsin: görselli sunum (PDF veya PowerPoint), belge/rapor (PDF veya Word) ve tablo (Excel). "
    "Kullanıcı bir sunum, slayt, PDF, rapor, teklif, belge, Word, Excel veya tablo dosyası isterse uzun içeriği sohbete "
    "yazma; 1-2 cümleyle ne hazırladığını söyle ve en sona tek satır ekle: [[DOSYA: <sunum|belge|tablo>|<pdf|pptx|docx|xlsx>"
    "|<içeriğin ayrıntılı açıklaması: konu, hedef kitle, slayt/bölüm sayısı, istenen her şey>]]. Format belirtilmemişse "
    "sunum ve belge için pdf, tablo için xlsx seç. Dosyayı sistem oluşturur ve indirme bağlantısını gösterir; dosya "
    "oluşturamayacağını asla söyleme."
)
FORCE_FILE_NOTE = "\n- Kullanıcı bu mesajda 'Dosya oluştur' seçeneğini kullandı: mutlaka en sona [[DOSYA: ...]] satırını ekle."

ACCENT = (79, 70, 229)  # indigo
DARK = (17, 17, 24)
GREY = (90, 90, 105)


class DocumentError(Exception):
    pass


def split_file_request(text: str) -> tuple[str, dict | None]:
    """[[DOSYA: tür|format|açıklama]] satırını ayırır: (görünen metin, istek veya None)."""
    m = FILE_TAG.search(text or "")
    if not m:
        return text, None
    parts = [p.strip() for p in m.group(1).split("|")]
    kind = KINDS.get((parts[0] if parts else "").lower(), "presentation")
    fmt = (parts[1] if len(parts) > 1 else "").lower().lstrip(".")
    if fmt in ("ppt", "powerpoint"):
        fmt = "pptx"
    if fmt in ("word", "doc"):
        fmt = "docx"
    if fmt in ("excel", "xls"):
        fmt = "xlsx"
    if fmt not in FORMATS[kind]:
        fmt = FORMATS[kind][0]
    brief = "|".join(parts[2:]).strip() if len(parts) > 2 else ""
    return FILE_TAG.sub("", text).strip(), {"kind": kind, "format": fmt, "brief": brief[:2000]}


# ---------------- İçerik (JSON) ----------------

SCHEMAS = {
    "presentation": (
        '{"title": "Sunum başlığı", "subtitle": "kısa alt başlık", "slides": [{"title": "Slayt başlığı", '
        '"bullets": ["en fazla 5 kısa madde"], "image_query": "2-4 English keywords for a stock photo", '
        '"note": "isteğe bağlı kısa konuşmacı notu"}]}',
        "6-10 slayt hazırla. Her slaytta 3-5 kısa, somut madde olsun (her biri en fazla 110 karakter). Son slayt özet "
        "veya çağrı (CTA) olsun.",
    ),
    "document": (
        '{"title": "Belge başlığı", "subtitle": "isteğe bağlı alt başlık", "sections": [{"heading": "Bölüm başlığı", '
        '"paragraphs": ["paragraf"], "bullets": ["isteğe bağlı maddeler"]}]}',
        "Kullanıcının istediği uzunlukta, düzenli bölümlerle, profesyonel ve somut bir belge hazırla.",
    ),
    "spreadsheet": (
        '{"title": "Tablo başlığı", "sheets": [{"name": "Sayfa adı (en fazla 30 karakter)", "columns": ["Sütun"], '
        '"rows": [["hücre değerleri; sayılar sayı olarak"]]}]}',
        "Gerçekçi ve kullanışlı satırlar oluştur; sayısal değerleri sayı olarak yaz.",
    ),
}


def _parse_json(text: str) -> dict:
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        raise DocumentError("İçerik JSON değil.")
    return json.loads(text[start : end + 1])


def build_spec(request: dict, conversation: list[dict], paid: bool = False) -> dict:
    """Sohbet bağlamından dosya içeriğini JSON olarak hazırlatır."""
    schema, guide = SCHEMAS[request["kind"]]
    context = "\n".join(
        f"{'Kullanıcı' if m['role'] == 'user' else 'Asistan'}: {m['content'] if isinstance(m['content'], str) else ''}"[:2500]
        for m in conversation[-8:]
        if m["role"] in ("user", "assistant")
    )
    prompt = (
        f"Sohbet:\n{context}\n\nİstenen dosya: {request['brief'] or 'yukarıdaki sohbete göre'}\n\n"
        f"Bugünün tarihi: {date.today():%d.%m.%Y}. {guide} Kullanıcının dilinde yaz (genelde Türkçe). "
        f"Yalnızca şu şemaya uyan geçerli JSON döndür, başka hiçbir şey yazma:\n{schema}"
    )
    last_error = None
    for _ in range(2):
        text, _ = smart_chat_completion([{"role": "user", "content": prompt}], max_tokens=6000, paid=paid)
        try:
            spec = _parse_json(text)
            if spec.get("title"):
                return spec
        except (ValueError, DocumentError) as e:
            last_error = e
    raise DocumentError(f"Dosya içeriği hazırlanamadı: {last_error}")


# ---------------- Stok görseller (sunum) ----------------

def _photo(query: str) -> bytes | None:
    from app.services import free_tier

    if not query or not free_tier.stock_enabled():
        return None
    try:
        if free_tier.stock_provider() == "pixabay":
            hits = free_tier._pixabay({"q": query[:100], "per_page": 3, "orientation": "horizontal"}).get("hits") or []
            url = hits[0].get("webformatURL") if hits else None
        else:
            photos = free_tier._pexels("search", {"query": query, "per_page": 1, "orientation": "landscape"}).get("photos") or []
            url = (photos[0].get("src") or {}).get("large") if photos else None
        if not url:
            return None
        r = requests.get(url, timeout=20)
        return r.content if r.status_code == 200 and len(r.content) < 8 * 1024 * 1024 else None
    except Exception as e:  # görsel olmadan devam
        logger.info("Slayt görseli alınamadı (%s): %s", query, e)
        return None


def _jpeg(data: bytes, max_side: int = 1400) -> bytes:
    from PIL import Image, ImageOps

    img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
    img.thumbnail((max_side, max_side))
    out = io.BytesIO()
    img.save(out, "JPEG", quality=85)
    return out.getvalue()


def _slide_images(slides: list[dict]) -> list[bytes | None]:
    from concurrent.futures import ThreadPoolExecutor

    with ThreadPoolExecutor(max_workers=6) as pool:
        raw = list(pool.map(lambda s: _photo(str(s.get("image_query") or "")), slides))
    out = []
    for data in raw:
        try:
            out.append(_jpeg(data) if data else None)
        except Exception:
            out.append(None)
    return out


# ---------------- Oluşturucular ----------------

def _str_list(v) -> list[str]:
    if isinstance(v, str):
        return [v]
    return [str(x) for x in (v or []) if str(x).strip()]


def _pdf():
    from fpdf import FPDF

    class PDF(FPDF):
        pass

    return PDF


def _fonts(pdf):
    pdf.add_font("Sans", "", str(FONT_DIR / "DejaVuSans.ttf"))
    pdf.add_font("Sans", "B", str(FONT_DIR / "DejaVuSans-Bold.ttf"))


def _cover(data: bytes, ratio: float) -> bytes:
    """Görseli verilen en/boy oranına ortadan kırpar (alanı tam doldurmak için)."""
    from PIL import Image

    img = Image.open(io.BytesIO(data)).convert("RGB")
    w, h = img.size
    if w / h > ratio:
        nw = int(h * ratio)
        img = img.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
    else:
        nh = int(w / ratio)
        img = img.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))
    out = io.BytesIO()
    img.save(out, "JPEG", quality=85)
    return out.getvalue()


def presentation_pdf(spec: dict) -> bytes:
    from fpdf import FPDF

    slides = spec.get("slides") or []
    images = _slide_images(slides)
    W, H = 338.7, 190.5  # 16:9 (13.33 x 7.5 inç)
    pdf = FPDF(orientation="L", unit="mm", format=(H, W))
    pdf.set_auto_page_break(False)
    pdf.set_margins(0, 0, 0)
    _fonts(pdf)
    LN = {"new_x": "LMARGIN", "new_y": "NEXT", "align": "L"}

    # Kapak: solda başlık, sağda ilk slaydın görseli
    pdf.add_page()
    pdf.set_fill_color(*DARK)
    pdf.rect(0, 0, W, H, "F")
    cover_img = next((i for i in images if i), None)
    text_w = W * 0.55 if cover_img else W - 56
    if cover_img:
        x0 = W * 0.6
        pdf.image(io.BytesIO(_cover(cover_img, (W - x0) / H)), x=x0, y=0, w=W - x0, h=H)
    pdf.set_fill_color(*ACCENT)
    pdf.rect(28, 58, 16, 16, "F")
    pdf.set_left_margin(28)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Sans", "B", 38)
    pdf.set_xy(28, 84)
    pdf.multi_cell(text_w - 28, 16, str(spec.get("title", "")), **LN)
    if spec.get("subtitle"):
        pdf.ln(4)
        pdf.set_font("Sans", "", 16)
        pdf.set_text_color(190, 190, 205)
        pdf.multi_cell(text_w - 28, 8, str(spec["subtitle"]), **LN)

    for i, (s, img) in enumerate(zip(slides, images), start=1):
        pdf.add_page()
        pdf.set_fill_color(255, 255, 255)
        pdf.rect(0, 0, W, H, "F")
        text_w = W * 0.55 if img else W - 60
        if img:
            x0 = W * 0.58
            pdf.image(io.BytesIO(_cover(img, (W - x0) / H)), x=x0, y=0, w=W - x0, h=H)
        pdf.set_fill_color(*ACCENT)
        pdf.rect(28, 26, 10, 3, "F")
        pdf.set_left_margin(28)
        pdf.set_text_color(*DARK)
        pdf.set_font("Sans", "B", 28)
        pdf.set_xy(28, 34)
        pdf.multi_cell(text_w - 34, 12.5, str(s.get("title", "")), **LN)
        pdf.set_font("Sans", "", 16)
        y = pdf.get_y() + 10
        for b in _str_list(s.get("bullets"))[:6]:
            pdf.set_fill_color(*ACCENT)
            pdf.rect(29, y + 3, 3, 3, "F")
            pdf.set_left_margin(38)
            pdf.set_xy(38, y)
            pdf.set_text_color(55, 55, 70)
            pdf.multi_cell(text_w - 44, 8, b, **LN)
            y = pdf.get_y() + 5
        pdf.set_font("Sans", "", 10)
        pdf.set_text_color(170, 170, 180)
        pdf.set_left_margin(0)
        pdf.set_xy(28, H - 16)
        pdf.cell(20, 5, f"{i:02d}")
    return bytes(pdf.output())


def presentation_pptx(spec: dict) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    slides = spec.get("slides") or []
    images = _slide_images(slides)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def box(slide, x, y, w, h, color):
        shp = slide.shapes.add_shape(1, x, y, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(*color)
        shp.line.fill.background()
        return shp

    def text(slide, x, y, w, h, value, size, bold=False, color=DARK):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = value
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor(*color)
        return tf

    s0 = prs.slides.add_slide(blank)
    box(s0, 0, 0, prs.slide_width, prs.slide_height, DARK)
    cover_img = next((i for i in images if i), None)
    if cover_img:
        x0 = Inches(8.0)
        w = prs.slide_width - x0
        s0.shapes.add_picture(io.BytesIO(_cover(cover_img, w / prs.slide_height)), x0, 0, w, prs.slide_height)
    box(s0, Inches(0.9), Inches(2.6), Inches(0.5), Inches(0.5), ACCENT)
    cw = Inches(6.8) if cover_img else Inches(11.5)
    text(s0, Inches(0.9), Inches(3.3), cw, Inches(1.6), str(spec.get("title", "")), 40, True, (255, 255, 255))
    if spec.get("subtitle"):
        text(s0, Inches(0.9), Inches(5.0), cw, Inches(1), str(spec["subtitle"]), 18, False, (190, 190, 205))

    for s, img in zip(slides, images):
        sl = prs.slides.add_slide(blank)
        text_w = Inches(6.6) if img else Inches(11.8)
        box(sl, Inches(0.9), Inches(0.95), Inches(0.45), Inches(0.12), ACCENT)
        text(sl, Inches(0.85), Inches(1.15), text_w, Inches(1.4), str(s.get("title", "")), 34, True)
        tb = sl.shapes.add_textbox(Inches(0.85), Inches(2.55), text_w, Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True
        for k, b in enumerate(_str_list(s.get("bullets"))[:6]):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.text = f"•  {b}"
            p.font.size = Pt(20)
            p.font.color.rgb = RGBColor(*GREY)
            p.space_after = Pt(10)
        if img:
            x0 = Inches(7.75)
            w = prs.slide_width - x0
            sl.shapes.add_picture(io.BytesIO(_cover(img, w / prs.slide_height)), x0, 0, w, prs.slide_height)
        if s.get("note"):
            sl.notes_slide.notes_text_frame.text = str(s["note"])
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def document_pdf(spec: dict) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_margins(22, 22, 22)
    pdf.set_auto_page_break(True, margin=20)
    _fonts(pdf)
    LN = {"new_x": "LMARGIN", "new_y": "NEXT", "align": "L"}
    pdf.add_page()
    pdf.set_fill_color(*ACCENT)
    pdf.rect(22, 22, 9, 9, "F")
    pdf.set_y(38)
    pdf.set_font("Sans", "B", 22)
    pdf.set_text_color(*DARK)
    pdf.multi_cell(0, 10, str(spec.get("title", "")), **LN)
    if spec.get("subtitle"):
        pdf.ln(1)
        pdf.set_font("Sans", "", 11.5)
        pdf.set_text_color(*GREY)
        pdf.multi_cell(0, 6, str(spec["subtitle"]), **LN)
    pdf.ln(6)
    for sec in spec.get("sections") or []:
        if sec.get("heading"):
            pdf.ln(2)
            pdf.set_font("Sans", "B", 14)
            pdf.set_text_color(*DARK)
            pdf.multi_cell(0, 8, str(sec["heading"]), **LN)
            pdf.ln(1.5)
        pdf.set_font("Sans", "", 11)
        pdf.set_text_color(50, 50, 60)
        for para in _str_list(sec.get("paragraphs")):
            pdf.multi_cell(0, 6.2, para, **LN)
            pdf.ln(2.5)
        for b in _str_list(sec.get("bullets")):
            y = pdf.get_y()
            pdf.set_fill_color(*ACCENT)
            pdf.rect(24, y + 2.3, 1.8, 1.8, "F")
            pdf.set_x(29)
            pdf.multi_cell(0, 6.2, b, **LN)
            pdf.ln(0.8)
        pdf.ln(2)
    return bytes(pdf.output())


def document_docx(spec: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    h = doc.add_heading(str(spec.get("title", "")), level=0)
    for run in h.runs:
        run.font.color.rgb = RGBColor(*DARK)
    if spec.get("subtitle"):
        doc.add_paragraph(str(spec["subtitle"])).runs[0].italic = True
    for sec in spec.get("sections") or []:
        if sec.get("heading"):
            doc.add_heading(str(sec["heading"]), level=1)
        for para in _str_list(sec.get("paragraphs")):
            doc.add_paragraph(para)
        for b in _str_list(sec.get("bullets")):
            doc.add_paragraph(b, style="List Bullet")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def spreadsheet_xlsx(spec: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for i, sh in enumerate(spec.get("sheets") or [{"name": "Sayfa1", "columns": [], "rows": []}]):
        name = re.sub(r"[\[\]:*?/\\]", "", str(sh.get("name") or f"Sayfa{i + 1}"))[:30] or f"Sayfa{i + 1}"
        ws = wb.create_sheet(name)
        cols = _str_list(sh.get("columns"))
        if cols:
            ws.append(cols)
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="4F46E5")
        for row in sh.get("rows") or []:
            ws.append(list(row) if isinstance(row, (list, tuple)) else [row])
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 50)
        ws.freeze_panes = "A2" if cols else None
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


BUILDERS = {
    ("presentation", "pdf"): presentation_pdf,
    ("presentation", "pptx"): presentation_pptx,
    ("document", "pdf"): document_pdf,
    ("document", "docx"): document_docx,
    ("spreadsheet", "xlsx"): spreadsheet_xlsx,
}


def _filename(title: str, fmt: str) -> str:
    tr = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")
    slug = re.sub(r"[^A-Za-z0-9]+", "-", (title or "ajansai").translate(tr)).strip("-").lower()[:60] or "ajansai"
    return f"{slug}.{fmt}"


def create(request: dict, conversation: list[dict], paid: bool = False) -> tuple[bytes, str, str, str]:
    """(dosya baytları, mime, dosya adı, başlık) döndürür."""
    spec = build_spec(request, conversation, paid)
    try:
        data = BUILDERS[(request["kind"], request["format"])](spec)
    except LLMError:
        raise
    except Exception as e:
        logger.exception("Dosya oluşturulamadı")
        raise DocumentError(f"Dosya oluşturulamadı: {e}") from e
    title = str(spec.get("title") or "Dosya")
    return data, MIME[request["format"]], _filename(title, request["format"]), title
