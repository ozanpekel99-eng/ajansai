"""Tasarımlı sunum motoru (PDF ve PowerPoint).

Claude her slayt için bir düzen (layout) seçer ve içeriği JSON olarak verir; burada aynı
çizim kodu iki arka uca (fpdf2 → PDF, python-pptx → PPTX) uygulanır. Yazılar Pillow ile
ölçülüp satırlara bölünür, böylece iki çıktı da aynı yerleşimi kullanır.

Düzenler: cover, cards, spotlight, table, split, stats, steps, bullets, quote, closing.
Vurgu: metinde *yıldız içindeki* kelimeler vurgu rengiyle yazılır.
"""

from __future__ import annotations

import io
import re
from functools import lru_cache
from pathlib import Path

FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
FONTS = {
    "display": ("Oswald", "Oswald-SemiBold.ttf"),
    "displayB": ("Oswald", "Oswald-Bold.ttf"),
    "body": ("Inter", "Inter-Regular.ttf"),
    "bodyB": ("Inter", "Inter-SemiBold.ttf"),
}
W, H = 960.0, 540.0
MX = 64.0  # yatay kenar boşluğu


def _hex(c: str) -> tuple[int, int, int]:
    c = c.lstrip("#")
    return int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)


PALETTES = {
    "gece": {"dark": "#141A23", "dsurf": "#1D2530", "dline": "#2E3846", "light": "#F1ECE3", "lsurf": "#FBF8F3",
             "lline": "#DCD3C5", "accent": "#FF5A1F", "accent2": "#F5C542"},
    "indigo": {"dark": "#12122A", "dsurf": "#1C1C3A", "dline": "#2E2E52", "light": "#F0EFF8", "lsurf": "#FAFAFE",
               "lline": "#D9D7EC", "accent": "#6C63FF", "accent2": "#FFB347"},
    "okyanus": {"dark": "#0E1B2C", "dsurf": "#16263B", "dline": "#25384F", "light": "#EDF3F7", "lsurf": "#F9FBFD",
                "lline": "#D2DEE7", "accent": "#16B3A6", "accent2": "#FFD166"},
    "orman": {"dark": "#13201A", "dsurf": "#1B2C23", "dline": "#2A3F33", "light": "#EEF0E6", "lsurf": "#FAFBF5",
              "lline": "#D5DAC6", "accent": "#E39B2D", "accent2": "#A8D5A2"},
    "bordo": {"dark": "#25131A", "dsurf": "#321C24", "dline": "#472B35", "light": "#F5EEE8", "lsurf": "#FCF8F5",
              "lline": "#E2D3C8", "accent": "#E0A042", "accent2": "#F08A8A"},
    "mono": {"dark": "#111111", "dsurf": "#1B1B1B", "dline": "#2D2D2D", "light": "#F2F2F0", "lsurf": "#FBFBFA",
             "lline": "#D8D8D4", "accent": "#E63946", "accent2": "#F4D35E"},
}
DEFAULT_PALETTE = "gece"


class Theme:
    def __init__(self, spec_theme: dict | None):
        spec_theme = spec_theme if isinstance(spec_theme, dict) else {}
        pal = dict(PALETTES.get(str(spec_theme.get("palette", "")).lower(), PALETTES[DEFAULT_PALETTE]))
        for key in ("accent", "accent2"):
            v = str(spec_theme.get(key) or "")
            if re.fullmatch(r"#?[0-9a-fA-F]{6}", v):
                pal[key] = "#" + v.lstrip("#")
        self.c = {k: _hex(v) for k, v in pal.items()}

    def tone(self, dark: bool) -> dict:
        c = self.c
        if dark:
            return {"bg": c["dark"], "surf": c["dsurf"], "line": c["dline"], "text": (245, 241, 234),
                    "muted": (160, 168, 180), "accent": c["accent"], "accent2": c["accent2"], "dark": True}
        return {"bg": c["light"], "surf": c["lsurf"], "line": c["lline"], "text": c["dark"],
                "muted": (95, 100, 112), "accent": c["accent"], "accent2": c["accent"], "dark": False}


# ---------------- Yazı ölçümü ----------------

@lru_cache(maxsize=64)
def _pil_font(key: str, size: float):
    from PIL import ImageFont

    return ImageFont.truetype(str(FONT_DIR / FONTS[key][1]), max(1, int(round(size * 4))))


def text_width(text: str, key: str, size: float, tracking: float = 0.0) -> float:
    if not text:
        return 0.0
    return _pil_font(key, size).getlength(text) / 4 + tracking * max(len(text) - 1, 0)


def tr_upper(s: str) -> str:
    return s.replace("i", "İ").replace("ı", "I").upper()


def runs_of(text: str) -> list[tuple[str, bool]]:
    """'Ev ve *kira hakları*' → [("Ev ve ", False), ("kira hakları", True)]"""
    out, accent = [], False
    for part in re.split(r"(\*)", str(text or "")):
        if part == "*":
            accent = not accent
        elif part:
            out.append((part, accent))
    return out


def plain(text) -> str:
    return str(text or "").replace("*", "").strip()


def wrap_runs(runs: list[tuple[str, bool]], key: str, size: float, width: float, tracking: float = 0.0):
    """Kelimelere bölüp satırlara dizer. Her satır [(metin, vurgu)] listesi."""
    words: list[tuple[str, bool]] = []
    for text, acc in runs:
        for i, w in enumerate(re.split(r"(\s+)", text)):
            if w:
                words.append((w, acc))
    lines, cur, cur_w = [], [], 0.0
    for w, acc in words:
        if w.isspace():
            if cur:
                cur.append((" ", acc))
                cur_w += text_width(" ", key, size, tracking)
            continue
        ww = text_width(w, key, size, tracking)
        if cur and cur_w + ww > width:
            while cur and cur[-1][0] == " ":
                cur.pop()
            lines.append(cur)
            cur, cur_w = [], 0.0
        cur.append((w, acc))
        cur_w += ww
    while cur and cur[-1][0] == " ":
        cur.pop()
    if cur:
        lines.append(cur)
    # aynı vurgudaki komşu parçaları birleştir
    merged = []
    for line in lines:
        m: list[list] = []
        for t, a in line:
            if m and m[-1][1] == a:
                m[-1][0] += t
            else:
                m.append([t, a])
        merged.append([(t, a) for t, a in m])
    return merged


def fit(text, key: str, max_size: float, min_size: float, width: float, max_lines: int,
        upper: bool = False, tracking: float = 0.0):
    """En büyük uygun punto: (punto, satırlar). Sığmazsa en küçük puntoda kırpar."""
    s = tr_upper(str(text or "")) if upper else str(text or "")
    runs = runs_of(s)
    size = max_size
    while True:
        lines = wrap_runs(runs, key, size, width, tracking)
        if len(lines) <= max_lines or size <= min_size:
            if len(lines) > max_lines:
                lines = lines[:max_lines]
                t, a = lines[-1][-1]
                lines[-1][-1] = (t.rstrip(" ,.;:") + "…", a)
            return size, lines
        size = max(min_size, size - 1)


# ---------------- Çizim arka uçları ----------------

class PdfCanvas:
    def __init__(self):
        from fpdf import FPDF

        self.pdf = FPDF(orientation="L", unit="pt", format=(H, W))
        self.pdf.set_auto_page_break(False)
        self.pdf.set_margins(0, 0, 0)
        for key, (_, file) in FONTS.items():
            self.pdf.add_font(key, "", str(FONT_DIR / file))

    def page(self, bg):
        self.pdf.add_page()
        self.rect(0, 0, W, H, bg)

    def rect(self, x, y, w, h, fill=None, radius=0.0, stroke=None, stroke_w=1.0):
        p = self.pdf
        style = ("F" if fill else "") + ("D" if stroke else "")
        if not style:
            return
        if fill:
            p.set_fill_color(*fill)
        if stroke:
            p.set_draw_color(*stroke)
            p.set_line_width(stroke_w)
        if radius:
            p.rect(x, y, w, h, style=style, round_corners=True, corner_radius=radius)
        else:
            p.rect(x, y, w, h, style=style)

    def ellipse(self, x, y, w, h, fill):
        self.pdf.set_fill_color(*fill)
        self.pdf.ellipse(x, y, w, h, style="F")

    def line(self, x1, y1, x2, y2, color, width=1.0):
        self.pdf.set_draw_color(*color)
        self.pdf.set_line_width(width)
        self.pdf.line(x1, y1, x2, y2)

    def lines(self, x, y, w, lines, key, size, color, accent, line_h=1.2, align="L", tracking=0.0):
        """Satırları çizer; son satırın altını (y) döndürür."""
        p = self.pdf
        p.set_font(key, size=size)
        p.set_char_spacing(tracking)
        font = _pil_font(key, size)
        asc, desc = (m / 4 for m in font.getmetrics())
        lh = size * line_h
        for i, line in enumerate(lines):
            lw = sum(text_width(t, key, size, tracking) for t, _ in line)
            cx = x + (w - lw if align == "R" else (w - lw) / 2 if align == "C" else 0)
            base = y + i * lh + (lh - (asc + desc)) / 2 + asc
            for t, acc in line:
                p.set_text_color(*(accent if acc else color))
                p.text(cx, base, t)
                cx += text_width(t, key, size, tracking)
        p.set_char_spacing(0)
        return y + len(lines) * lh

    def image(self, data: bytes, x, y, w, h):
        self.pdf.image(io.BytesIO(data), x=x, y=y, w=w, h=h)

    def output(self) -> bytes:
        return bytes(self.pdf.output())


class PptxCanvas:
    def __init__(self):
        from pptx import Presentation
        from pptx.util import Pt

        self.Pt = Pt
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(W), Pt(H)
        self.blank = self.prs.slide_layouts[6]
        self.slide = None

    def _rgb(self, c):
        from pptx.dml.color import RGBColor

        return RGBColor(*c)

    def page(self, bg):
        self.slide = self.prs.slides.add_slide(self.blank)
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb(bg)

    def rect(self, x, y, w, h, fill=None, radius=0.0, stroke=None, stroke_w=1.0):
        from pptx.enum.shapes import MSO_SHAPE

        Pt = self.Pt
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)
        )
        if radius:
            shape.adjustments[0] = min(0.5, radius / max(min(w, h), 1))
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = self._rgb(fill)
        else:
            shape.fill.background()
        if stroke:
            shape.line.color.rgb = self._rgb(stroke)
            shape.line.width = Pt(stroke_w)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def ellipse(self, x, y, w, h, fill):
        from pptx.enum.shapes import MSO_SHAPE

        Pt = self.Pt
        shape = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(x), Pt(y), Pt(w), Pt(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._rgb(fill)
        shape.line.fill.background()
        shape.shadow.inherit = False

    def line(self, x1, y1, x2, y2, color, width=1.0):
        if y1 == y2:
            self.rect(x1, y1 - width / 2, x2 - x1, width, color)
        else:
            self.rect(x1 - width / 2, y1, width, y2 - y1, color)

    def lines(self, x, y, w, lines, key, size, color, accent, line_h=1.2, align="L", tracking=0.0):
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

        Pt = self.Pt
        if not lines:
            return y
        lh = size * line_h
        tb = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w + 4), Pt(lh * len(lines)))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        family = FONTS[key][0]
        bold = key in ("displayB", "bodyB")
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"L": PP_ALIGN.LEFT, "C": PP_ALIGN.CENTER, "R": PP_ALIGN.RIGHT}[align]
            para.line_spacing = Pt(lh)
            for t, acc in line:
                r = para.add_run()
                r.text = t
                f = r.font
                f.name = family
                f.size = Pt(size)
                f.bold = bold
                f.color.rgb = self._rgb(accent if acc else color)
                if tracking:
                    r._r.get_or_add_rPr().set("spc", str(int(tracking * 100)))
        return y + len(lines) * lh

    def image(self, data: bytes, x, y, w, h):
        Pt = self.Pt
        self.slide.shapes.add_picture(io.BytesIO(data), Pt(x), Pt(y), Pt(w), Pt(h))

    def notes(self, text: str):
        if text:
            self.slide.notes_slide.notes_text_frame.text = text

    def output(self) -> bytes:
        out = io.BytesIO()
        self.prs.save(out)
        return out.getvalue()


# ---------------- Yardımcı çizimler ----------------

def _cover_crop(data: bytes, w: float, h: float) -> bytes:
    from PIL import Image, ImageOps

    img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
    ratio = w / h
    iw, ih = img.size
    if iw / ih > ratio:
        nw = int(ih * ratio)
        img = img.crop(((iw - nw) // 2, 0, (iw - nw) // 2 + nw, ih))
    else:
        nh = int(iw / ratio)
        img = img.crop((0, (ih - nh) // 2, iw, (ih - nh) // 2 + nh))
    img.thumbnail((1600, 1600))
    out = io.BytesIO()
    img.save(out, "JPEG", quality=86)
    return out.getvalue()


def _items(v, limit: int) -> list:
    return [x for x in (v if isinstance(v, list) else []) if x][:limit]


def _txt(d, key) -> str:
    return str((d or {}).get(key) or "").strip() if isinstance(d, dict) else str(d or "").strip()


def kicker(cv, t, x, y, text, w=W - 2 * MX):
    if not text:
        return y
    size, lines = fit(text, "bodyB", 9.5, 8, w, 1, upper=True, tracking=1.4)
    return cv.lines(x, y, w, lines, "bodyB", size, t["accent"], t["accent"], tracking=1.4) + 6


def heading(cv, t, x, y, w, text, max_size=40, min_size=24, max_lines=2):
    size, lines = fit(text, "display", max_size, min_size, w, max_lines, upper=True)
    return cv.lines(x, y, w, lines, "display", size, t["text"], t["accent2"] if t["dark"] else t["accent"], line_h=1.08)


def para(cv, t, x, y, w, text, size=12, max_lines=6, color=None, key="body", min_size=9):
    if not text:
        return y
    s, lines = fit(text, key, size, min_size, w, max_lines)
    return cv.lines(x, y, w, lines, key, s, color or t["muted"], t["accent"], line_h=1.42)


def card_box(cv, t, x, y, w, h):
    cv.rect(x, y, w, h, t["surf"], radius=8, stroke=t["line"], stroke_w=0.8)


def pill(cv, x, y, text, bg, fg, size=8.5):
    text = tr_upper(text)
    tw = text_width(text, "bodyB", size, 0.6)
    cv.rect(x, y, tw + 20, size + 11, bg, radius=(size + 11) / 2)
    cv.lines(x + 10, y + 5.5, tw + 2, [[(text, False)]], "bodyB", size, fg, fg, line_h=1.0, tracking=0.6)
    return tw + 20


# ---------------- Düzenler ----------------

def l_cover(cv, t, s):
    y = 150 if s.get("tiles") else 175
    y = kicker(cv, t, MX, y - 40, _txt(s, "kicker"))
    size, lines = fit(s.get("title"), "displayB", 64, 36, W - 2 * MX - 80, 3, upper=True)
    y = cv.lines(MX, y + 4, W - 2 * MX, lines, "displayB", size, t["text"], t["accent2"], line_h=1.02)
    y = para(cv, t, MX, y + 14, 640, _txt(s, "subtitle"), size=14, max_lines=2)
    tiles = _items(s.get("tiles"), 6)
    if tiles:
        gap = 10
        tw = (W - 2 * MX - gap * (len(tiles) - 1)) / len(tiles)
        ty = max(y + 40, 380)
        for i, tile in enumerate(tiles):
            tx = MX + i * (tw + gap)
            cv.rect(tx, ty, tw, 92, t["surf"], radius=4)
            cv.rect(tx, ty, tw, 3, t["accent"])
            num = _txt(tile, "num") or f"{i + 1:02d}"
            cv.lines(tx + 14, ty + 14, tw - 28, [[(num, False)]], "displayB", 20, t["accent"], t["accent"], line_h=1.0)
            para(cv, t, tx + 14, ty + 44, tw - 28, _txt(tile, "label"), size=10, max_lines=3, color=t["text"], min_size=8)
    elif s.get("note"):
        para(cv, t, MX, H - 70, 600, _txt(s, "note"), size=10, max_lines=1)


def l_cards(cv, t, s):
    y = kicker(cv, t, MX, 58, _txt(s, "kicker"))
    y = heading(cv, t, MX, y, W - 2 * MX, s.get("title"), 34, 24, 2) + 24
    cards = _items(s.get("cards"), 4)
    bar = s.get("bar") if isinstance(s.get("bar"), dict) else None
    n = max(len(cards), 1)
    gap = 14
    cw = (W - 2 * MX - gap * (n - 1)) / n
    ch = (H - 64 - y) - (78 if bar else 0)
    ch = min(ch, 220)
    for i, c in enumerate(cards):
        cx = MX + i * (cw + gap)
        card_box(cv, t, cx, y, cw, ch)
        yy = y + 18
        if _txt(c, "label"):
            yy = cv.lines(cx + 18, yy, cw - 36, fit(_txt(c, "label"), "bodyB", 8.5, 7.5, cw - 36, 1, upper=True, tracking=0.8)[1],
                          "bodyB", 8.5, t["accent"], t["accent"], tracking=0.8) + 8
        size, lines = fit(_txt(c, "title"), "display", 20, 14, cw - 36, 2, upper=True)
        yy = cv.lines(cx + 18, yy, cw - 36, lines, "display", size, t["text"], t["accent"], line_h=1.1) + 8
        para(cv, t, cx + 18, yy, cw - 36, _txt(c, "text"), size=10.5, max_lines=int((y + ch - yy - 14) / 15), min_size=8.5)
    if bar:
        by = y + ch + 20
        cv.rect(MX, by, W - 2 * MX, 54, t["text"] if not t["dark"] else t["surf"], radius=6)
        fg = t["bg"] if not t["dark"] else t["text"]
        label = tr_upper(_txt(bar, "label"))
        cv.lines(MX + 18, by + 19, 120, [[(label, False)]], "displayB", 13, t["accent2"] if not t["dark"] else t["accent"], t["accent"], line_h=1.2)
        items = _items(bar.get("items"), 4)
        if items:
            iw = (W - 2 * MX - 160) / len(items)
            for i, it in enumerate(items):
                ix = MX + 150 + i * iw
                cv.ellipse(ix, by + 24, 6, 6, t["accent"])
                _, ls = fit(str(it), "body", 9.5, 8, iw - 24, 2)
                cv.lines(ix + 12, by + (19 if len(ls) == 1 else 13), iw - 24, ls, "body", 9.5, fg, t["accent"], line_h=1.35)


def _phone(cv, t, x, y, w, h, statement, top, badge):
    cv.rect(x, y, w, h, (12, 15, 20), radius=22)
    cv.rect(x + 6, y + 6, w - 12, h - 12, (24, 29, 37) if True else t["surf"], radius=17)
    if top:
        cv.lines(x + 20, y + 22, w - 40, fit(top, "body", 8, 7, w - 40, 1)[1], "body", 8, (150, 156, 168), t["accent"], line_h=1.2)
    size, lines = fit(statement, "displayB", 22, 13, w - 40, 7, upper=True)
    block = len(lines) * size * 1.05
    sy = y + (h - block) / 2 - 6
    cv.lines(x + 20, sy, w - 40, lines, "displayB", size, (245, 241, 234), t["accent2"], line_h=1.05)
    if badge:
        pill(cv, x + 20, y + h - 46, badge, t["accent"], (255, 255, 255), size=7.5)


def l_spotlight(cv, t, s):
    pw, ph = 200, 360
    px, py = MX + 10, (H - ph) / 2
    if str(s.get("frame") or "phone") == "panel":
        cv.rect(px, py, pw, ph, t["surf"], radius=10, stroke=t["accent"], stroke_w=1.5)
        size, lines = fit(s.get("statement"), "displayB", 26, 14, pw - 40, 8, upper=True)
        cv.lines(px + 20, py + (ph - len(lines) * size * 1.05) / 2, pw - 40, lines, "displayB", size, t["text"], t["accent"], line_h=1.05)
    else:
        _phone(cv, t, px, py, pw, ph, s.get("statement"), _txt(s, "top"), _txt(s, "badge"))
    x = px + pw + 56
    w = W - MX - x
    y = kicker(cv, t, x, 92, _txt(s, "kicker"), w)
    y = heading(cv, t, x, y, w, s.get("title"), 36, 22, 2) + 10
    cv.line(x, y, x + w, y, t["line"], 0.8)
    y += 18
    items = _items(s.get("items"), 4)
    avail = H - 70 - y
    per = avail / max(len(items), 1)
    for i, it in enumerate(items):
        iy = y + i * per
        cv.ellipse(x, iy + 3, 9, 9, t["accent"])
        yy = cv.lines(x + 22, iy, w - 22, fit(_txt(it, "title"), "bodyB", 12, 10, w - 22, 1)[1], "bodyB", 12, t["text"], t["accent"], line_h=1.3)
        para(cv, t, x + 22, yy + 3, w - 22, _txt(it, "text"), size=10.5, max_lines=max(1, int((per - 24) / 15)), min_size=8.5)


def l_table(cv, t, s):
    y = kicker(cv, t, MX, 54, _txt(s, "kicker"))
    y = heading(cv, t, MX, y, W - 2 * MX, s.get("title"), 30, 20, 1) + 18
    cols = [str(c) for c in _items(s.get("columns"), 5)]
    rows = [r if isinstance(r, list) else [r] for r in _items(s.get("rows"), 9)]
    if not cols:
        return
    n = len(cols)
    tw = W - 2 * MX
    # sütun genişlikleri: içerik uzunluğuna göre
    weights = []
    for j in range(n):
        longest = max([len(str(cols[j]))] + [len(str(r[j])) if j < len(r) else 0 for r in rows])
        weights.append(min(max(longest, 6), 70))
    total = sum(weights)
    widths = [tw * w_ / total for w_ in weights]
    acc_col = s.get("accent_col")
    acc_col = acc_col if isinstance(acc_col, int) and 0 <= acc_col < n else (n - 1 if n > 2 else None)
    cv.rect(MX, y, tw, 24, t["surf"] if t["dark"] else t["text"])
    hx = MX
    hdr_color = t["accent"] if t["dark"] else t["bg"]
    for j, c in enumerate(cols):
        cv.lines(hx + 10, y + 7.5, widths[j] - 20, fit(c, "bodyB", 8.5, 7.5, widths[j] - 20, 1)[1], "bodyB", 8.5, hdr_color, hdr_color)
        hx += widths[j]
    y += 24
    avail = H - 50 - y
    font = 9.5
    rendered = []
    for r in rows:
        cells = []
        for j in range(n):
            val = str(r[j]) if j < len(r) else ""
            key = "bodyB" if j == acc_col else "body"
            _, ls = fit(val, key, font, font, widths[j] - 20, 4)
            cells.append((ls, key))
        rendered.append(cells)
    heights = [max(len(c[0]) for c in cells) * font * 1.35 + 14 for cells in rendered]
    scale = min(1.0, avail / max(sum(heights), 1))
    for cells, rh in zip(rendered, heights):
        rh *= scale
        cx = MX
        for j, (ls, key) in enumerate(cells):
            color = t["accent2"] if (j == acc_col and t["dark"]) else (t["accent"] if j == acc_col else (t["muted"] if j == 0 else t["text"]))
            cv.lines(cx + 10, y + 7, widths[j] - 20, ls, key, font * min(1, scale + 0.15), color, t["accent"], line_h=1.35)
            cx += widths[j]
        y += rh
        cv.line(MX, y, MX + tw, y, t["line"], 0.6)


def l_split(cv, t, s, image: bytes | None):
    has_img = bool(image)
    w = 420 if has_img else W - 2 * MX
    y = kicker(cv, t, MX, 78, _txt(s, "kicker"), w)
    y = heading(cv, t, MX, y, w, s.get("title"), 38, 22, 3) + 16
    y = para(cv, t, MX, y, w, _txt(s, "text"), size=12, max_lines=4) + 14
    for b in _items(s.get("bullets"), 5):
        cv.rect(MX, y + 6, 10, 2.5, t["accent"])
        _, ls = fit(str(b), "body", 11.5, 9.5, w - 22, 3)
        y = cv.lines(MX + 22, y, w - 22, ls, "body", 11.5, t["text"], t["accent"], line_h=1.4) + 9
    if has_img:
        x0 = MX + w + 40
        cv.image(_cover_crop(image, W - x0, H), x0, 0, W - x0, H)
        cv.rect(x0, 0, 5, H, t["accent"])


def l_stats(cv, t, s):
    y = kicker(cv, t, MX, 70, _txt(s, "kicker"))
    y = heading(cv, t, MX, y, W - 2 * MX, s.get("title"), 36, 22, 2) + 40
    stats = _items(s.get("stats"), 4)
    n = max(len(stats), 1)
    gap = 28
    cw = (W - 2 * MX - gap * (n - 1)) / n
    for i, st in enumerate(stats):
        x = MX + i * (cw + gap)
        cv.rect(x, y, cw, 3, t["accent"])
        size, ls = fit(_txt(st, "value"), "displayB", 60, 30, cw, 1)
        yy = cv.lines(x, y + 18, cw, ls, "displayB", size, t["accent2"] if t["dark"] else t["accent"], t["accent"], line_h=1.05) + 6
        yy = cv.lines(x, yy, cw, fit(_txt(st, "label"), "bodyB", 12, 10, cw, 2)[1], "bodyB", 12, t["text"], t["accent"], line_h=1.3) + 6
        para(cv, t, x, yy, cw, _txt(st, "text"), size=10.5, max_lines=5)


def l_steps(cv, t, s):
    y = kicker(cv, t, MX, 64, _txt(s, "kicker"))
    y = heading(cv, t, MX, y, W - 2 * MX, s.get("title"), 34, 22, 2) + 44
    steps = _items(s.get("steps"), 5)
    n = max(len(steps), 1)
    gap = 18
    cw = (W - 2 * MX - gap * (n - 1)) / n
    cv.line(MX + 14, y + 14, W - MX - cw + 14, y + 14, t["line"], 1.2)
    for i, st in enumerate(steps):
        x = MX + i * (cw + gap)
        cv.ellipse(x, y, 28, 28, t["accent"])
        cv.lines(x, y + 7, 28, [[(str(i + 1), False)]], "displayB", 12, (255, 255, 255), (255, 255, 255), line_h=1.1, align="C")
        yy = y + 46
        if _txt(st, "label"):
            yy = cv.lines(x, yy, cw, fit(_txt(st, "label"), "bodyB", 8.5, 7.5, cw, 1, upper=True, tracking=0.8)[1], "bodyB", 8.5,
                          t["accent"], t["accent"], tracking=0.8) + 6
        size, ls = fit(_txt(st, "title"), "display", 18, 13, cw, 2, upper=True)
        yy = cv.lines(x, yy, cw, ls, "display", size, t["text"], t["accent"], line_h=1.1) + 8
        para(cv, t, x, yy, cw, _txt(st, "text"), size=10.5, max_lines=8, min_size=8.5)


def l_bullets(cv, t, s):
    y = kicker(cv, t, MX, 64, _txt(s, "kicker"))
    y = heading(cv, t, MX, y, W - 2 * MX, s.get("title"), 34, 22, 2) + 28
    items = _items(s.get("bullets"), 8)
    cols = 2 if len(items) > 4 else 1
    cw = (W - 2 * MX - 40 * (cols - 1)) / cols
    per_col = -(-len(items) // cols)
    for c in range(cols):
        yy = y
        for i, b in enumerate(items[c * per_col:(c + 1) * per_col]):
            num = f"{c * per_col + i + 1:02d}"
            cv.lines(MX + c * (cw + 40), yy, 34, [[(num, False)]], "displayB", 16, t["accent"], t["accent"], line_h=1.2)
            title, _, rest = str(b).partition(":")
            x = MX + c * (cw + 40) + 40
            if rest.strip():
                yy2 = cv.lines(x, yy + 1, cw - 40, fit(title, "bodyB", 12, 10, cw - 40, 1)[1], "bodyB", 12, t["text"], t["accent"], line_h=1.35)
                yy2 = para(cv, t, x, yy2 + 2, cw - 40, rest.strip(), size=10.5, max_lines=3, min_size=9)
            else:
                yy2 = cv.lines(x, yy + 1, cw - 40, fit(b, "body", 12, 10, cw - 40, 3)[1], "body", 12, t["text"], t["accent"], line_h=1.4)
            yy = yy2 + 16
            cv.line(x, yy - 8, x + cw - 40, yy - 8, t["line"], 0.6)


def l_quote(cv, t, s):
    cv.lines(MX, 96, 80, [[("“", False)]], "displayB", 110, t["accent"], t["accent"], line_h=1.0)
    size, lines = fit(s.get("statement") or s.get("title"), "display", 40, 22, W - 2 * MX - 60, 4, upper=True)
    y = cv.lines(MX + 60, 190, W - 2 * MX - 60, lines, "display", size, t["text"], t["accent2"] if t["dark"] else t["accent"], line_h=1.1)
    if _txt(s, "author"):
        cv.rect(MX + 60, y + 24, 30, 2.5, t["accent"])
        para(cv, t, MX + 100, y + 17, 500, _txt(s, "author"), size=11, max_lines=1)


def l_closing(cv, t, s):
    y = kicker(cv, t, MX, 150, _txt(s, "kicker"))
    size, lines = fit(s.get("title"), "displayB", 56, 30, W - 2 * MX - 60, 3, upper=True)
    y = cv.lines(MX, y + 4, W - 2 * MX, lines, "displayB", size, t["text"], t["accent2"] if t["dark"] else t["accent"], line_h=1.03)
    y = para(cv, t, MX, y + 16, 620, _txt(s, "text"), size=13, max_lines=3) + 26
    if _txt(s, "cta"):
        pill(cv, MX, y, _txt(s, "cta"), t["accent"], (255, 255, 255), size=10)
    if _txt(s, "contact"):
        para(cv, t, MX, H - 72, W - 2 * MX, _txt(s, "contact"), size=10.5, max_lines=1)


LAYOUTS = {
    "cover": l_cover, "cards": l_cards, "spotlight": l_spotlight, "table": l_table, "stats": l_stats,
    "steps": l_steps, "bullets": l_bullets, "quote": l_quote, "closing": l_closing,
}
DARK_BY_DEFAULT = {"cover", "table", "quote", "closing"}


def _normalize(spec: dict) -> list[dict]:
    """Eski biçim (başlık + maddeler) ve eksik alanlar için güvenli slayt listesi."""
    slides = [s for s in (spec.get("slides") or []) if isinstance(s, dict)]
    out = []
    if not slides or str(slides[0].get("layout")) != "cover":
        out.append({"layout": "cover", "title": spec.get("title") or "Sunum", "subtitle": spec.get("subtitle"),
                    "kicker": spec.get("kicker")})
    for s in slides:
        layout = str(s.get("layout") or "").lower()
        if layout not in LAYOUTS and layout != "split":
            layout = "split" if s.get("image_query") else "bullets"
        out.append({**s, "layout": layout})
    return out


def _tone_for(i: int, s: dict, prev_dark: bool | None) -> bool:
    tone = str(s.get("tone") or "").lower()
    if tone in ("dark", "koyu"):
        return True
    if tone in ("light", "acik", "açık"):
        return False
    if s["layout"] in DARK_BY_DEFAULT or i == 0:
        return True
    return not prev_dark if prev_dark is not None else False


def render(spec: dict, cv, images: dict[int, bytes]) -> bytes:
    theme = Theme(spec.get("theme"))
    prev = None
    for i, s in enumerate(_normalize(spec)):
        dark = _tone_for(i, s, prev)
        prev = dark
        t = theme.tone(dark)
        cv.page(t["bg"])
        if s["layout"] == "split":
            l_split(cv, t, s, images.get(i))
        else:
            LAYOUTS[s["layout"]](cv, t, s)
        if hasattr(cv, "notes"):
            cv.notes(_txt(s, "note"))
    return cv.output()


def image_queries(spec: dict) -> dict[int, str]:
    """Fotoğraf gereken slaytlar (split düzeni): {slayt sırası: arama sorgusu}."""
    return {i: str(s.get("image_query")) for i, s in enumerate(_normalize(spec))
            if s["layout"] == "split" and s.get("image_query")}


def to_pdf(spec: dict, images: dict[int, bytes]) -> bytes:
    return render(spec, PdfCanvas(), images)


def to_pptx(spec: dict, images: dict[int, bytes]) -> bytes:
    return render(spec, PptxCanvas(), images)


SCHEMA = """{
  "title": "Sunum başlığı",
  "theme": {"palette": "gece | indigo | okyanus | orman | bordo | mono", "accent": "#RRGGBB (isteğe bağlı, marka rengi)"},
  "slides": [
    {"layout": "cover", "kicker": "MARKA · KONU · TARİH", "title": "İki satırlık güçlü başlık, *vurgulu kısım*", "subtitle": "tek cümlelik açıklama", "tiles": [{"num": "01", "label": "kısa etiket"}]},
    {"layout": "cards", "kicker": "...", "title": "...", "cards": [{"label": "0-3 sn / ADIM 1 / ÖZELLİK", "title": "KISA BAŞLIK", "text": "1-2 cümle"}], "bar": {"label": "OFİS SETUP", "items": ["kısa madde"]}},
    {"layout": "spotlight", "frame": "phone | panel", "top": "@hesap · Reels", "statement": "Büyük vurucu cümle, *vurgulu kısım*", "badge": "0-3 sn · Hook", "kicker": "...", "title": "...", "items": [{"title": "Alt başlık", "text": "1-2 cümle"}]},
    {"layout": "table", "kicker": "...", "title": "...", "columns": ["Süre", "Ne söylenir", "Ekran yazısı"], "rows": [["0-3 sn", "metin", "EKRAN YAZISI"]], "accent_col": 2},
    {"layout": "split", "kicker": "...", "title": "...", "text": "kısa paragraf", "bullets": ["madde"], "image_query": "2-4 English keywords for a stock photo"},
    {"layout": "stats", "kicker": "...", "title": "...", "stats": [{"value": "%38", "label": "etiket", "text": "kısa açıklama"}]},
    {"layout": "steps", "kicker": "...", "title": "...", "steps": [{"label": "1. HAFTA", "title": "Adım başlığı", "text": "1-2 cümle"}]},
    {"layout": "bullets", "kicker": "...", "title": "...", "bullets": ["Başlık: açıklama"]},
    {"layout": "quote", "statement": "Akılda kalan tek cümle, *vurgu*", "author": "kaynak"},
    {"layout": "closing", "kicker": "...", "title": "Kapanış başlığı, *vurgu*", "text": "kısa özet", "cta": "Kısa çağrı", "contact": "iletişim / hesap"}
  ]
}"""

GUIDE = (
    "Profesyonel bir tasarım ajansı kalitesinde, içeriği dolu ve somut bir sunum hazırla. Kurallar: "
    "1) İlk slayt cover, son slayt closing olsun. 2) Arka arkaya aynı düzeni kullanma; içeriğe en uygun düzeni seç "
    "(süreç → steps, karşılaştırma/özellikler → cards, sayılar → stats, senaryo/metin/plan → table, tek fikir + örnek → "
    "spotlight, fotoğrafın anlam kattığı yer → split; split'i en fazla 2 slaytta kullan). 3) Başlıklar kısa ve vurucu "
    "(en fazla ~8 kelime); önemli kısmı *yıldız* içine al. 4) Kicker kısa üst etiket: konu · bölüm · süre gibi. "
    "5) Metinler gerçek, uygulanabilir ve kullanıcının işine özel olsun; genel geçer laf yazma, örnek, sayı ve somut "
    "detay ver. 6) Kart metni en fazla 2 kısa cümle, tablo hücresi en fazla ~25 kelime. 7) Paleti konuya göre seç; "
    "marka rengi biliniyorsa accent olarak ver. 8) Kullanıcı slayt sayısı söylemediyse 7-12 slayt."
)
