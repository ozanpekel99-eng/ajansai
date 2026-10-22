import io
import json

from app import config
from app.services import documents, free_tier, storage

from .conftest import auth
from .test_api import BLOB_TOKEN, BLOB_URL, _enable_media, register_paid

PRES = {
    "title": "Yaz Kampanyası",
    "subtitle": "Çorlu Kafe · Ekim 2026",
    "slides": [
        {"title": "Hedefler", "bullets": ["Instagram takipçisini %20 artırmak", "Soğuk kahve satışını ikiye katlamak"], "image_query": "iced coffee"},
        {"title": "Takvim", "bullets": ["1. hafta: tanıtım", "2. hafta: çekiliş"], "image_query": "calendar"},
    ],
}
DOC = {"title": "Teklif", "sections": [{"heading": "Kapsam", "paragraphs": ["Şişli'de çekim, ğüşıöç."], "bullets": ["Reels", "Gönderi"]}]}
SHEET = {"title": "Bütçe", "sheets": [{"name": "Ekim", "columns": ["Kalem", "Tutar"], "rows": [["Reklam", 5000], ["Çekim", 3000]]}]}


def _no_photos(monkeypatch):
    monkeypatch.setattr(documents, "_photo", lambda q: None)


def test_split_file_request():
    text, req = documents.split_file_request("Hazırlıyorum.\n[[DOSYA: sunum|PowerPoint|Kafe için 8 slaytlık sunum]]")
    assert text == "Hazırlıyorum." and req == {"kind": "presentation", "format": "pptx", "brief": "Kafe için 8 slaytlık sunum"}
    assert documents.split_file_request("[[DOSYA: tablo|pdf|bütçe]]")[1]["format"] == "xlsx"
    assert documents.split_file_request("düz metin") == ("düz metin", None)


def test_builders_produce_valid_files(monkeypatch):
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (800, 500), "orange").save(buf, "JPEG")
    monkeypatch.setattr(documents, "_photo", lambda q: buf.getvalue())

    pdf = documents.presentation_pdf(PRES)
    assert pdf.startswith(b"%PDF") and len(pdf) > 2000
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) == 3 and "Kampanyası" in reader.pages[0].extract_text()

    from pptx import Presentation

    prs = Presentation(io.BytesIO(documents.presentation_pptx(PRES)))
    assert len(prs.slides) == 3
    assert any(sh.shape_type == 13 for sh in prs.slides[1].shapes)  # resim var

    doc_pdf = documents.document_pdf(DOC)
    assert "ğüşıöç" in PdfReader(io.BytesIO(doc_pdf)).pages[0].extract_text()

    from docx import Document

    d = Document(io.BytesIO(documents.document_docx(DOC)))
    assert any("Kapsam" in p.text for p in d.paragraphs)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(documents.spreadsheet_xlsx(SHEET)))
    assert wb["Ekim"]["B2"].value == 5000


def test_free_chat_creates_presentation_file(client, monkeypatch):
    _no_photos(monkeypatch)
    monkeypatch.setattr(config, "BLOB_READ_WRITE_TOKEN", BLOB_TOKEN)
    monkeypatch.setattr(free_tier, "free_chat_completion", lambda messages, max_tokens=0, vision=False: ("Sunumunu hazırladım.\n[[DOSYA: sunum|pdf|Kafe yaz kampanyası]]", "groq:x"))
    monkeypatch.setattr(documents, "free_chat_completion", lambda messages, max_tokens=0: ("```json\n" + json.dumps(PRES) + "\n```", "groq:x"))
    saved = {}
    monkeypatch.setattr(storage, "save", lambda d, m, f, filename=None: saved.update(data=d, mime=m, name=filename) or BLOB_URL + filename)
    r = client.post("/api/free/chat", json={"mode": "file", "messages": [{"role": "user", "content": "Kafem için sunum hazırla"}]}, headers={"x-forwarded-for": "7.7.7.7"})
    data = r.json()
    assert r.status_code == 200, r.text
    assert data["reply"] == "Sunumunu hazırladım." and data["files"][0]["name"] == "yaz-kampanyasi.pdf"
    assert saved["mime"] == "application/pdf" and saved["data"].startswith(b"%PDF")


def test_free_file_daily_limit(client, monkeypatch):
    _no_photos(monkeypatch)
    monkeypatch.setattr(free_tier, "FREE_FILES_PER_DAY", 1)
    monkeypatch.setattr(free_tier, "free_chat_completion", lambda messages, max_tokens=0, vision=False: ("Tamam.\n[[DOSYA: tablo|xlsx|bütçe]]", "groq:x"))
    monkeypatch.setattr(documents, "free_chat_completion", lambda messages, max_tokens=0: (json.dumps(SHEET), "groq:x"))
    ip = {"x-forwarded-for": "8.8.8.1"}
    body = {"messages": [{"role": "user", "content": "Bütçe tablosu"}]}
    assert client.post("/api/free/chat", json=body, headers=ip).json()["files"][0]["url"].startswith("data:application/vnd.openxml")
    second = client.post("/api/free/chat", json=body, headers=ip).json()
    assert second["files"] == [] and "dosya oluşturabilirsin" in second["image_error"]


def test_paid_chat_file_costs_two_credits_and_refunds_on_failure(client, monkeypatch):
    _enable_media(monkeypatch)
    _no_photos(monkeypatch)
    from app.services import assistant

    monkeypatch.setattr(assistant, "chat_completion", lambda model, messages, max_tokens=0: "Hazır.\n[[DOSYA: belge|docx|teklif]]")
    monkeypatch.setattr(documents, "free_chat_completion", lambda messages, max_tokens=0: (json.dumps(DOC), "groq:x"))
    monkeypatch.setattr(storage, "save", lambda d, m, f, filename=None: BLOB_URL + filename)
    token = register_paid(client)
    r = client.post("/api/content/chat", json={"messages": [{"role": "user", "content": "Teklif yaz"}]}, headers=auth(token))
    assert r.json()["credits_spent"] == 3 and r.json()["files"][0]["name"] == "teklif.docx"

    monkeypatch.setattr(documents, "free_chat_completion", lambda messages, max_tokens=0: ("bozuk", "groq:x"))
    r = client.post("/api/content/chat", json={"messages": [{"role": "user", "content": "Teklif yaz"}]}, headers=auth(token))
    assert r.json()["files"] == [] and r.json()["credits_spent"] == 1 and r.json()["image_error"]
    used = client.get("/api/plans/usage", headers=auth(token)).json()["credits"]["used"]
    assert used == 4  # 3 + 1; başarısız dosyanın 2 kredisi iade edildi
